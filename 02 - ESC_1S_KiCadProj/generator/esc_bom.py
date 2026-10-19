# ACwing ESC1S rev B: assembly BoM (+ Costing sheet) and JLCPCB BoM.
#   python3 esc_bom.py <out dir>
import sys, os, re, json
import pcbnew
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

OUT = sys.argv[1]
PROJECT = 'ACwing_ESC1S_B'
parts = json.load(open(os.path.join(OUT, PROJECT + '.kicad_sch.parts.json')))
b = pcbnew.LoadBoard(os.path.join(OUT, PROJECT + '.kicad_pcb'))
side = {f.GetReference(): ('Bottom' if f.IsFlipped() else 'Top') for f in b.GetFootprints()}
key = lambda r: (re.sub(r'\d', '', r), int(re.sub(r'\D', '', r) or 0))
Q, E = 'LCSC quote', 'Estimate'
# (value, footprint) -> MPN, manufacturer, LCSC/JLC #, package, description, note, JLC footprint label, unit price
#                       10 boards, unit price 100 boards (USD), price source
TAB = {
 ('AON7524', 'Vishay_PowerPAK_1212-8_Single'): ('AON7524', 'Alpha & Omega Semiconductor', 'C431195', 'DFN 3x3 EP (DFN3x3A_8L_EP1)',
     'MOSFET, N-channel, 30 V, 30 A, logic level', 'Low-side motor switch, one per channel. Land pattern: Vishay PowerPAK 1212-8 (check against the AOS drawing)', 'DFN-8 3x3', 0.232, 0.178, Q),
 ('STPS3L45AF', 'D_SMAflat'): ('SS34F', 'GOODWORK', 'C908669', 'SMAF (DO-221AC)', 'Diode, Schottky, 40 V, 3 A, Vf 0.55 V @ 3 A',
     'Motor freewheel diode. Replaces STPS3L45AF (low LCSC stock), same DO-221AC outline. Cathode to VBAT', 'SMAF', 0.0326, 0.0326, Q),
 ('TSC213', 'SOT-363_SC-70-6'): ('TSC213ICT', 'STMicroelectronics', 'C1858319', 'SC70-6 (SOT-363)', 'Current-sense amplifier, gain 50 V/V, zero drift',
     'Battery current: CUR = 0.1 V/A on J1 pin 3', 'SC-70-6', 0.60, 0.45, E),
 ('2m', 'R_1206_3216Metric'): ('RLM12FTCMR002', 'TA-I Tech', 'C710257', '1206 (3216 metric)', 'Resistor, current sense, 2 mOhm, 1 %, 1 W',
     'Battery shunt, bottom side, Kelvin-connected to U1', '1206', 0.0476, 0.0389, Q),
 ('470uF 6.3V polymer', 'CP_EIA-7343-31_Kemet-D'): ('T520D477M006ATE015', 'KEMET', 'C696809', '7343-31 (EIA D), h 2.8 mm',
     'Capacitor, tantalum polymer, 470 uF, 6.3 V, 15 mOhm', 'Bulk at the battery entry, + to VBAT. Low LCSC stock: any 470 uF 6.3 V 7343 polymer fits', '7343-31', 1.6307, 1.2652, Q),
 ('10uF 10V', 'C_0603_1608Metric'): ('CL10A106KP8NNNC', 'Samsung Electro-Mechanics', 'C19702', '0603 (1608 metric)', 'Capacitor, MLCC, 10 uF, 10 V, X5R',
     'Local decoupling next to each MOSFET (hot loop)', '0603', 0.0287, 0.0287, Q),
 ('100nF', 'C_0402_1005Metric'): ('CL05B104KO5NNNC', 'Samsung Electro-Mechanics', 'C1525', '0402 (1005 metric)', 'Capacitor, MLCC, 100 nF, 16 V, X7R',
     'U1 supply (C7) and CUR filter (C6)', '0402', 0.0049, 0.0049, Q),
 ('22R', 'R_0402_1005Metric'): ('0402WGF220JTCE', 'UNI-ROYAL', 'C25092', '0402 (1005 metric)', 'Resistor, 22 Ohm, 5 %, 1/16 W', 'MOSFET gate resistors', '0402', 0.002, 0.002, E),
 ('100k', 'R_0402_1005Metric'): ('0402WGF1003TCE', 'UNI-ROYAL', 'C25741', '0402 (1005 metric)', 'Resistor, 100 kOhm, 1 %, 1/16 W',
     'Gate pull-downs: MOSFETs off with the cable unplugged or the FC booting (R5-R8); Q5 gate pull-down (R11)', '0402', 0.002, 0.002, E),
 ('1k', 'R_0402_1005Metric'): ('0402WGF1001TCE', 'UNI-ROYAL', 'C11702', '0402 (1005 metric)', 'Resistor, 1 kOhm, 1 %, 1/16 W', 'CUR filter (1k / 100 nF, 1.6 kHz) (R10); charge LED (R15)', '0402', 0.002, 0.002, E),
 ('ESC link', 'JST_SH_SM08B-SRSS-TB_1x08-1MP_P1.00mm_Horizontal'): ('SM08B-SRSS-TB(LF)(SN)', 'JST', 'C160407', 'JST SH 8P side entry, h 2.9 mm',
     'Connector, JST SH, 8 positions, 1.0 mm pitch, side entry', 'ESC link to the FC: 1 GND, 2 VBAT, 3 CUR, 4 TLM (n.c.), 5-8 M1-M4. Mouth at the rear edge', 'SMD,P=1mm,Side entry (JST SH 8P)', 0.35, 0.25, E),
 ('AON7423', 'Vishay_PowerPAK_1212-8_Single'): ('AON7423', 'Alpha & Omega Semiconductor', 'C99711', 'DFN 3x3 EP',
     'MOSFET, P-channel, 20 V, 50 A, 4.9 mOhm @ -2.5 V, Vgs +-8 V', 'Charge switch Q5: board supply off while the charge port is plugged in', 'DFN-8 3x3', 0.8993, 0.7663, Q),
 ('STC4054GR', 'TSOT-23-5'): ('STC4054GR', 'STMicroelectronics', 'C262930', 'TSOT-23-5', 'Li-ion linear charger, 4.2 V, up to 800 mA',
     '1S charger, 455 mA (R14 2.2k)', 'TSOT-23-5', 1.4469, 1.0963, Q),
 ('RB521S30', 'D_SOD-523'): ('RB521S30T1G', 'onsemi', 'C8523', 'SOD-523', 'Diode, Schottky, 30 V, 200 mA', 'VBUS -> Q5 gate (charge switch off while charging)', 'SOD-523', 0.03, 0.02, E),
 ('green', 'LED_0402_1005Metric'): ('XL-1005UGC', 'XINGLIGHT', 'C965793', '0402 (1005 metric)', 'LED, green 525 nm, 0402', 'Charge LED (on = charging), top side next to the charge port', '0402', 0.012, 0.008, E),
 ('USB-C (HRO TYPE-C-31-M-12)', 'USB_C_Receptacle_HRO_TYPE-C-31-M-12'): ('TYPE-C-31-M-12', 'HRO (Korea Hroparts)', 'C165948', 'USB-C receptacle, 16 pin, SMD + THT shell',
     'Connector, USB Type-C receptacle, top mount', 'Charge port (power only), front-left chamfer', 'SMD (USB-C 16P)', 0.1729, 0.1365, Q),
 ('5.1k', 'R_0402_1005Metric'): ('0402WGF5101TCE', 'UNI-ROYAL', 'C25905', '0402 (1005 metric)', 'Resistor, 5.1 kOhm, 1 %, 1/16 W', 'USB-C CC pull-downs (sink)', '0402', 0.002, 0.002, E),
 ('2.2k 1%', 'R_0402_1005Metric'): ('0402WGF2201TCE', 'UNI-ROYAL', 'C25879', '0402 (1005 metric)', 'Resistor, 2.2 kOhm, 1 %, 1/16 W', 'Charger R_PROG: 455 mA', '0402', 0.002, 0.002, E),
 ('1R', 'R_0402_1005Metric'): ('FRC0402F1R00TS', 'FOJAN', 'C2909329', '0402 (1005 metric)', 'Resistor, 1 Ohm, 1 %, 1/16 W', 'Charger output RC (with C9)', '0402', 0.003, 0.002, E),
 ('1uF', 'C_0402_1005Metric'): ('CL05A105KA5NQNC', 'Samsung Electro-Mechanics', 'C52923', '0402 (1005 metric)', 'Capacitor, MLCC, 1 uF, 25 V, X5R', 'Charger input (VBUS)', '0402', 0.0068, 0.0068, Q),
 ('4.7uF 10V', 'C_0603_1608Metric'): ('CL10A475KO8NNNC', 'Samsung Electro-Mechanics', 'C19666', '0603 (1608 metric)', 'Capacitor, MLCC, 4.7 uF, 16 V, X5R', 'Charger output RC (with R16)', '0603', 0.0205, 0.0159, Q),
}
NOT_ASSEMBLED = [
 ('H', 'M2', 'N/A (PCB feature)', 'N/A', 'NPTH 2.2 mm', 'Mounting hole, M2, 20 x 20 pattern', 'Hard mount to the frame bosses; keep-out r 2.6 mm on both sides'),
 ('M', 'coreless 8520', 'N/A (PCB feature)', 'N/A', 'MotorPads_1x02_SMD', 'Motor lead solder pads (8.5 x 20 mm coreless brushed motors, 3.7 V)',
  'Soldered by hand; "+" pad (silk mark) to the red/blue lead. M1/M4 CW, M2/M3 CCW'),
 ('J', 'BATT 1S', 'N/A (PCB feature)', 'N/A', 'BattPads_1x02_SMD', 'Battery lead solder pads, 1S LiPo, AWG 22-24', 'Rear corner; "+" silk mark on the VBAT pad'),
 ('PIGTAIL', 'JST PH 2.0 pigtail', 'PH 2.0 male lead, 30-40 mm, 24 AWG (any brand)', 'N/A', 'off-board', 'Battery plug pigtail soldered to J2 (rated 2 A)', 'Hand-soldered: red to +. Match the polarity of the battery plug (not standardised)'),
]
groups, skipped, dnp = {}, [], []
for ref, v in parts.items():
    fpn = v['fp'].split(':')[1]
    if ref.startswith(('TP', 'H', 'M')) or ref == 'J2':
        skipped.append(ref); continue
    if v.get('dnp'):
        dnp.append(ref); continue
    k = (v['value'], fpn)
    if k not in TAB:
        raise SystemExit('no MPN for %s %s' % (ref, k))
    groups.setdefault((k, side[ref]), []).append(ref)

FONT = 'Arial'
thin = Side(style='thin', color='BFBFBF')
def style_header(ws, r):
    for c in ws[r]:
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='305496')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
wb = Workbook(); ws = wb.active; ws.title = 'BoM'
hdr = ['Item', 'Quantity', 'Part Reference', 'Value', 'MPN', 'Manufacturer', 'LCSC / JLCPCB #', 'Package', 'Description', 'Side', 'Notes']
ws.append(['ACwing ESC1S rev B - assembly BoM (1 board)'])
ws.append(['Source: ACwing_ESC1S_B.kicad_sch / ACwing_ESC1S_B.kicad_pcb. LCSC part numbers from the rev E FC BoM (2026-10-15/17); J1 C160407, rev B parts (Q5 AON7423 C99711, U2 STC4054GR C262930, J4 TYPE-C-31-M-12 C165948) checked 2026-10-18. D1-D4: SS34F fitted in place of STPS3L45AF (same outline).'])
ws.append([]); ws.append(hdr); H = 4
rows = sorted(groups.items(), key=lambda kv: key(sorted(kv[1], key=key)[0]))
for i, ((k, sd), refs) in enumerate(rows, 1):
    t = TAB[k]; refs = sorted(refs, key=key)
    ws.append([i, len(refs), ', '.join(refs), k[0], t[0], t[1], t[2], t[3], t[4], sd, t[5]])
last = ws.max_row
ws.append([])
ws.append(['Total', '=SUM(B%d:B%d)' % (H + 1, last), 'Total placements (assembled parts)', '-', '=COUNTA(E%d:E%d)' % (H + 1, last),
           'line items', '-', '-', '-', '-', 'Not-assembled items listed below'])
tot = ws.max_row
ws.append([]); ws.append(['Not assembled (PCB features / DNP)']); na0 = ws.max_row
ws.append(hdr); NA_H = ws.max_row
n = 0
for pref, val, mpn, man, pkg, desc, note in NOT_ASSEMBLED:
    if pref == 'PIGTAIL':
        n += 1
        ws.append(['NA%d' % n, 1, '- (soldered to J2)', val, mpn, man, 'N/A', pkg, desc, 'Top', note]); continue
    refs = sorted([r for r in skipped if re.match(pref + r'\d', r)] if pref != 'J' else ['J2'], key=key)
    n += 1
    ws.append(['NA%d' % n, len(refs), ', '.join(refs), val, mpn, man, 'N/A', pkg, desc, ', '.join(sorted({side[r] for r in refs})), note])
for r in sorted(dnp, key=key):
    n += 1
    ws.append(['NA%d' % n, 1, r, parts[r]['value'], '0402WGF4701TCE', 'UNI-ROYAL', 'C25900 (DNP)', '0402 (1005 metric)',
               'Resistor, 4.7 kOhm, 1 %, 1/16 W', side[r], 'Do not populate: the I2C pull-ups are on the FC (R15/R16). Fit only for a host without pull-ups'])
end = ws.max_row
for row in ws.iter_rows(min_row=1, max_row=end):
    for c in row:
        c.font = Font(name=FONT, size=10); c.alignment = Alignment(wrap_text=True, vertical='top')
ws['A1'].font = Font(name=FONT, size=13, bold=True)
ws['A2'].font = Font(name=FONT, size=9, italic=True, color='595959')
style_header(ws, H); style_header(ws, NA_H)
ws.cell(na0, 1).font = Font(name=FONT, size=11, bold=True)
for r in list(range(H + 1, last + 1)) + list(range(NA_H + 1, end + 1)):
    for c in ws[r]:
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for c in ws[tot]:
    c.font = Font(name=FONT, size=10, bold=True)
for col, w in zip('ABCDEFGHIJK', (6, 9, 22, 14, 24, 22, 20, 24, 40, 9, 60)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A5'

# ---------------------------------------------------------------- Costing
cs = wb.create_sheet('Costing')
cs.append(['ACwing ESC1S rev B - component costing (assembled parts, USD)'])
cs.append(['Blue = input prices; black = formulas. "Estimate" rows are typical distributor prices, not quotes. PCB (4 layers), assembly fees, motors and the 8-pin cable are not included.'])
cs.append([])
ch = ['References', 'Qty / board', 'MPN', 'Source', 'Unit price @10 boards', 'Unit price @100 boards', 'Cost / board @10', 'Cost / board @100', 'Price basis']
cs.append(ch); CH = 4
for (k, sd), refs in rows:
    t = TAB[k]; r = cs.max_row + 1
    cs.append([', '.join(sorted(refs, key=key)), len(refs), t[0], t[2], t[7], t[8], '=B%d*E%d' % (r, r), '=B%d*F%d' % (r, r), t[9]])
cl = cs.max_row
cs.append([])
cs.append(['Total per board', '=SUM(B%d:B%d)' % (CH + 1, cl), '', '', '', '', '=SUM(G%d:G%d)' % (CH + 1, cl), '=SUM(H%d:H%d)' % (CH + 1, cl), ''])
ct = cs.max_row
for row in cs.iter_rows(min_row=1, max_row=cs.max_row):
    for c in row:
        c.font = Font(name=FONT, size=10); c.alignment = Alignment(wrap_text=True, vertical='top')
cs['A1'].font = Font(name=FONT, size=13, bold=True)
cs['A2'].font = Font(name=FONT, size=9, italic=True, color='595959')
style_header(cs, CH)
for r in range(CH + 1, cl + 1):
    for col in (5, 6):
        cs.cell(r, col).font = Font(name=FONT, size=10, color='0000FF'); cs.cell(r, col).number_format = '$#,##0.000'
    for col in (7, 8):
        cs.cell(r, col).number_format = '$#,##0.000'
    for c in cs[r]:
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for col in (7, 8):
    cs.cell(ct, col).number_format = '$#,##0.00'
for c in cs[ct]:
    c.font = Font(name=FONT, size=10, bold=True)
for col, w in zip('ABCDEFGHI', (22, 10, 24, 24, 13, 13, 14, 14, 16)):
    cs.column_dimensions[col].width = w
wb.save(os.path.join(OUT, PROJECT + '_BoM.xlsx'))

# ---------------------------------------------------------------- JLCPCB BoM
jw = Workbook(); js = jw.active; js.title = 'Sheet1'
js.append(['Comment', 'Designator', 'Footprint', 'JLCPCB Part #'])
jg = {}
for (k, sd), refs in rows:
    t = TAB[k]
    g = jg.setdefault(t[0], [t, []]); g[1] += refs
for mpn, (t, refs) in sorted(jg.items(), key=lambda kv: key(sorted(kv[1][1], key=key)[0])):
    js.append([mpn, ','.join(sorted(refs, key=key)), t[6], t[2] if t[2].startswith('C') and t[2][1:].isdigit() else ''])
for row in js.iter_rows():
    for c in row:
        c.font = Font(name=FONT, size=10, bold=(c.row == 1))
for col, w in zip('ABCD', (26, 30, 34, 16)):
    js.column_dimensions[col].width = w
jw.save(os.path.join(OUT, PROJECT + '_BoM_JLCPCB.xlsx'))
print('BoM lines', len(rows), 'placements', sum(len(v) for v in groups.values()), '| JLC lines', len(jg), '| not assembled', n)
