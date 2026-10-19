# Renames every DimaQuad project in the repository to ACwing (file names, library names, texts, spreadsheet cells).
import os, sys, subprocess, re
from openpyxl import load_workbook
ROOT = sys.argv[1]
REPL = [
 ('DimaQuadFC_', 'ACwing_FC_'), ('DimaQuadESC_', 'ACwing_ESC1S_'), ('DimaQuadGPS_', 'ACwing_GPS_'), ('DimaQuadFlow_', 'ACwing_Flow_'),
 ('DimaQuadGPS.pretty', 'ACwing_GPS.pretty'), ('DimaQuadFlow.pretty', 'ACwing_Flow.pretty'),
 ('DimaQuadGPS:', 'ACwing_GPS:'), ('DimaQuadFlow:', 'ACwing_Flow:'),
 ('"DimaQuadGPS"', '"ACwing_GPS"'), ('"DimaQuadFlow"', '"ACwing_Flow"'), ("'DimaQuadGPS'", "'ACwing_GPS'"), ("'DimaQuadFlow'", "'ACwing_Flow'"),
 ('DimaQuad 1S power board', 'ACwing 1S power board'), ('DimaQuad Stack Check', 'ACwing Stack Check'),
 ('DimaQuadFC', 'ACwing FC'), ('DimaQuadESC', 'ACwing ESC1S'), ('DimaQuadGPS', 'ACwing GPS'), ('DimaQuadFlow', 'ACwing Flow'),
 ('DimaMiniQuad', 'ACwing'), ('DimaQuad.pretty', 'ACwing.pretty'), ('DimaQuad.kicad_sym', 'ACwing.kicad_sym'),
 ('DimaQuad:', 'ACwing:'), ('"DimaQuad"', '"ACwing"'), ("'DimaQuad'", "'ACwing'"), ('DimaQuad_', 'ACwing_'), ('DimaQuad', 'ACwing'),
 ('UAV Control - A. Cirillo', 'ACwing - Andrea Cirillo'), ('A. Cirillo', 'Andrea Cirillo'),
]
def fix(s):
    for a, b in REPL:
        s = s.replace(a, b)
    return s
def fix_name(n):
    for a, b in REPL:
        if ' ' in b or b.startswith(('"', "'")):
            continue
        n = n.replace(a, b)
    return n
TEXT = ('.kicad_sch', '.kicad_pcb', '.kicad_pro', '.kicad_prl', '.kicad_sym', '.kicad_mod', '.kicad_dru', '.md', '.py', '.sh',
        '.json', '.rpt', '.html', '.js', '.dxf', 'fp-lib-table', 'sym-lib-table')
nt = nx = nr = 0
for dp, dns, fns in os.walk(ROOT):
    if '/.git' in dp:
        continue
    for fn in fns:
        p = os.path.join(dp, fn)
        if fn.endswith(TEXT):
            s = open(p, encoding='utf-8', errors='surrogateescape').read()
            s2 = fix(s)
            if s2 != s:
                open(p, 'w', encoding='utf-8', errors='surrogateescape').write(s2); nt += 1
        elif fn.endswith('.xlsx'):
            wb = load_workbook(p); ch = False
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value, str) and ('DimaQuad' in c.value or 'Cirillo' in c.value):
                            c.value = fix(c.value); ch = True
            if ch:
                wb.save(p); nx += 1
# renames, deepest first
for dp, dns, fns in sorted(os.walk(ROOT), key=lambda t: -t[0].count('/')):
    if '/.git' in dp:
        continue
    for n in fns + dns:
        m = fix_name(n)
        if m != n:
            subprocess.run(['git', '-C', ROOT, 'mv', os.path.join(dp, n), os.path.join(dp, m)], check=True); nr += 1
print('text files changed', nt, '| spreadsheets', nx, '| renamed', nr)
