# Release renaming: the current boards become rev A (identifiers in file names, texts and spreadsheet cells).
import os, sys, subprocess, re
from openpyxl import load_workbook
ROOT = sys.argv[1]
REPL = [('ACwing_FC_H', 'ACwing_FC_A'), ('ACwing_ESC1S_B', 'ACwing_ESC1S_A'), ('ACwing_GPS_D', 'ACwing_GPS_A'),
        ('ACwing FC rev H', 'ACwing FC rev A'), ('ACwing ESC1S rev B', 'ACwing ESC 1S rev A'), ('ACwing ESC 1S rev B', 'ACwing ESC 1S rev A'),
        ('ACwing GPS rev D', 'ACwing GPS rev A')]
def fix(s):
    for a, b in REPL: s = s.replace(a, b)
    return s
TEXT = ('.kicad_sch', '.kicad_pcb', '.kicad_pro', '.kicad_prl', '.kicad_sym', '.kicad_mod', '.kicad_dru', '.md', '.py', '.sh',
        '.json', '.rpt', '.html', '.js', '.dxf', 'fp-lib-table', 'sym-lib-table')
nt = nx = nr = 0
for dp, dns, fns in os.walk(ROOT):
    if '/.git' in dp: continue
    for fn in fns:
        p = os.path.join(dp, fn)
        if fn.endswith(TEXT):
            s = open(p, encoding='utf-8', errors='surrogateescape').read(); s2 = fix(s)
            if s2 != s: open(p, 'w', encoding='utf-8', errors='surrogateescape').write(s2); nt += 1
        elif fn.endswith('.xlsx'):
            wb = load_workbook(p); ch = False
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value, str):
                            v = fix(c.value)
                            v = re.sub(r'top side \(rev D\)', 'top side', v)
                            v = v.replace('; rev E: on SPI1', '; draft E: on SPI1').replace('from the rev E FC BoM', 'from the draft E FC BoM')
                            if v != c.value: c.value = v; ch = True
            if ch: wb.save(p); nx += 1
for dp, dns, fns in sorted(os.walk(ROOT), key=lambda t: -t[0].count('/')):
    if '/.git' in dp: continue
    for n in fns + dns:
        m = n
        for a, b in REPL[:3]: m = m.replace(a, b)
        if m != n:
            subprocess.run(['git', '-C', ROOT, 'mv', os.path.join(dp, n), os.path.join(dp, m)], check=True); nr += 1
print('text files changed', nt, '| spreadsheets', nx, '| renamed', nr)
