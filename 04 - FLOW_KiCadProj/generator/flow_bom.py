# ACwing Flow rev A: assembly BoM (+ Costing sheet) and JLCPCB BoM.
#   python3 gps_bom.py <out dir>
import sys, os, re, json
import pcbnew
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

OUT = sys.argv[1]
PROJECT = 'ACwing_Flow_A'
parts = json.load(open(os.path.join(OUT, PROJECT + '.kicad_sch.parts.json')))
b = pcbnew.LoadBoard(os.path.join(OUT, PROJECT + '.kicad_pcb'))
side = {f.GetReference(): ('Bottom' if f.IsFlipped() else 'Top') for f in b.GetFootprints()}
key = lambda r: (re.sub(r'\d', '', r), int(re.sub(r'\D', '', r) or 0))
Q, E = 'LCSC quote', 'Estimate'
# (value, footprint) -> MPN, manufacturer, LCSC/JLC #, package, description, note, JLC footprint label, unit price
#                       10 boards, unit price 100 boards (USD), price source
TAB = {
 ('PMW3901MB-TXQT', 'PixArt_PMW3901MB_COB-28_6x6mm'): ('PMW3901MB-TXQT', 'PixArt Imaging', 'C43496881', 'COB-28 6 x 6 mm (lens 6 x 6 x 2.28)',
     'Optical flow sensor, far field (80 mm to infinity), SPI 2 MHz', 'Needs the LN03-ZSZ lens assembly: check whether the LCSC/JLC part includes it, otherwise order the lens separately and fit it after reflow', 'COB-28 6x6', 5.20, 4.60, E),
 ('VL53L1CXV0FY/1', 'ST_VL53L1x'): ('VL53L1CXV0FY/1', 'STMicroelectronics', 'C190004', 'LGA-12 4.9 x 2.5 x 1.56 mm',
     'Time-of-flight distance sensor, up to 4 m, I2C 0x29', 'Height above ground; keep the cover-glass area clear (no window)', 'LGA-12 4.9x2.5', 3.20, 2.70, E),
 ('STM32C011F4U6', 'ST_UFQFPN-20_3x3mm_P0.5mm'): ('STM32C011F4U6TR', 'STMicroelectronics', 'C5456188', 'UFQFPN-20 3 x 3 mm',
     'MCU Cortex-M0+, 48 MHz, 16 KB flash, 6 KB RAM', 'Flow bridge: reads the PMW3901 on SPI, I2C target 0x3B on the Qwiic bus', 'UFQFPN-20 (3x3)', 1.6331, 1.2305, Q),
 ('LDK130PU-R', 'DFN-6_1.3x1.2mm_P0.4mm'): ('LDK130PU-R', 'STMicroelectronics', 'C157446', 'DFN-6 1.2 x 1.3 mm',
     'LDO, adjustable, 300 mA', 'PMW3901 VDD 2.0 V (R5/R6)', 'DFN-6 (1.2x1.3)', 0.4186, 0.3181, Q),
 ('Qwiic', 'JST_SH_SM04B-SRSS-TB_1x04-1MP_P1.00mm_Horizontal'): ('SM04B-SRSS-TB(LF)(SN)', 'JST', 'C160404', 'JST SH 4P side entry, h 2.9 mm',
     'Connector, JST SH, 4 positions, 1.0 mm pitch, side entry', 'Qwiic in/out (pass-through), left and right edges', 'SMD,P=1mm,Side entry (JST SH 4P)', 0.25, 0.18, E),
 ('22uF', 'C_0603_1608Metric'): ('CL10A226MQ8NRNC', 'Samsung Electro-Mechanics', 'C59461', '0603 (1608 metric)', 'Capacitor, MLCC, 22 uF, 6.3 V, X5R', 'PMW3901 VDD (PixArt reference)', '0603', 0.012, 0.010, E),
 ('10uF', 'C_0603_1608Metric'): ('CL10A106KP8NNNC', 'Samsung Electro-Mechanics', 'C19702', '0603 (1608 metric)', 'Capacitor, MLCC, 10 uF, 10 V, X5R', 'PMW3901 VDDIO', '0603', 0.0287, 0.0287, Q),
 ('4.7uF', 'C_0603_1608Metric'): ('CL10A475KO8NNNC', 'Samsung Electro-Mechanics', 'C19666', '0603 (1608 metric)', 'Capacitor, MLCC, 4.7 uF, 16 V, X5R', 'PMW3901 VREG, VL53L1X and MCU bulk', '0603', 0.0205, 0.0159, Q),
 ('1uF', 'C_0402_1005Metric'): ('CL05A105KA5NQNC', 'Samsung Electro-Mechanics', 'C52923', '0402 (1005 metric)', 'Capacitor, MLCC, 1 uF, 25 V, X5R', 'LDO input / output', '0402', 0.0068, 0.0068, Q),
 ('100nF', 'C_0402_1005Metric'): ('CL05B104KO5NNNC', 'Samsung Electro-Mechanics', 'C1525', '0402 (1005 metric)', 'Capacitor, MLCC, 100 nF, 16 V, X7R', 'Decoupling', '0402', 0.0049, 0.0049, Q),
 ('green', 'LED_0402_1005Metric'): ('XL-1005UGC', 'XINGLIGHT', 'C965793', '0402 (1005 metric)', 'LED, green, 0402', 'Status LED (PA8)', '0402', 0.012, 0.008, E),
 ('3R', 'R_0402_1005Metric'): ('0402WGF300KTCE', 'UNI-ROYAL', 'C423196', '0402 (1005 metric)', 'Resistor, 3 Ohm, 1 %, 1/16 W', 'PMW3901 VDD filter (PixArt reference)', '0402', 0.002, 0.002, E),
 ('1k', 'R_0402_1005Metric'): ('0402WGF1001TCE', 'UNI-ROYAL', 'C11702', '0402 (1005 metric)', 'Resistor, 1 kOhm, 1 %, 1/16 W', 'LED current', '0402', 0.002, 0.002, E),
 ('10k', 'R_0402_1005Metric'): ('0402WGF1002TCE', 'UNI-ROYAL', 'C25744', '0402 (1005 metric)', 'Resistor, 10 kOhm, 1 %, 1/16 W', 'Pull-ups (NRESET, NCS, MOTION, XSHUT, INT)', '0402', 0.002, 0.002, E),
 ('15k 1%', 'R_0402_1005Metric'): ('0402WGF1502TCE', 'UNI-ROYAL', 'C25756', '0402 (1005 metric)', 'Resistor, 15 kOhm, 1 %, 1/16 W', 'LDO divider, top', '0402', 0.002, 0.002, E),
 ('10k 1%', 'R_0402_1005Metric'): ('0402WGF1002TCE', 'UNI-ROYAL', 'C25744', '0402 (1005 metric)', 'Resistor, 10 kOhm, 1 %, 1/16 W', 'LDO divider, bottom', '0402', 0.002, 0.002, E),
}
NOT_ASSEMBLED = [
 ('H', 'M2', 'N/A (PCB feature)', 'N/A', 'NPTH 2.2 mm', 'Mounting hole, M2, 16 x 16 pattern', 'Keep-out r 2.4 mm on both sides for the nylon screw / standoff'),
 ('TP', 'Test pad', 'N/A (PCB feature)', 'N/A', 'SMD pad 1.0 mm', 'SWD pads: SWDIO, SWCLK, NRST, 3.3 V, GND', 'First programming of the STM32C011 with a pogo fixture or tag-connect style probe'),
]
groups, skipped, dnp = {}, [], []
for ref, v in parts.items():
    fpn = v['fp'].split(':')[1]
    if ref.startswith(('TP', 'H')):
        skipped.append(ref); continue
    if v.get('dnp'):
        dnp.append(ref); continue
    k = (v['value'], fpn)
    if k not in TAB:
        raise SystemExit('no MPN for %s %s' % (ref, k))
    groups.setdefault((k, side[ref]), []).append(ref)

FONT = 'Arial'
thin = Side(style='thin', color='BFBFBF')
def style_header(ws, r):
    for c in ws[r]:
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='305496')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
wb = Workbook(); ws = wb.active; ws.title = 'BoM'
hdr = ['Item', 'Quantity', 'Part Reference', 'Value', 'MPN', 'Manufacturer', 'LCSC / JLCPCB #', 'Package', 'Description', 'Side', 'Notes']
ws.append(['ACwing Flow rev A - assembly BoM (1 board)'])
ws.append(['Source: ACwing_Flow_A.kicad_sch / ACwing_Flow_A.kicad_pcb. LCSC part numbers checked on 2026-10-17. Check the PMW3901 lens (LN03-ZSZ) supply before ordering.'])
ws.append([]); ws.append(hdr); H = 4
rows = sorted(groups.items(), key=lambda kv: key(sorted(kv[1], key=key)[0]))
for i, ((k, sd), refs) in enumerate(rows, 1):
    t = TAB[k]; refs = sorted(refs, key=key)
    ws.append([i, len(refs), ', '.join(refs), k[0], t[0], t[1], t[2], t[3], t[4], sd, t[5]])
last = ws.max_row
ws.append([])
ws.append(['Total', '=SUM(B%d:B%d)' % (H + 1, last), 'Total placements (assembled parts)', '-', '=COUNTA(E%d:E%d)' % (H + 1, last),
           'line items', '-', '-', '-', '-', 'Not-assembled items listed below'])
tot = ws.max_row
ws.append([]); ws.append(['Not assembled (PCB features / DNP)']); na0 = ws.max_row
ws.append(hdr); NA_H = ws.max_row
n = 0
for pref, val, mpn, man, pkg, desc, note in NOT_ASSEMBLED:
    refs = sorted([r for r in skipped if re.match(pref + r'\d', r)], key=key)
    n += 1
    ws.append(['NA%d' % n, len(refs), ', '.join(refs), val, mpn, man, 'N/A', pkg, desc, ', '.join(sorted({side[r] for r in refs})), note])
for r in sorted(dnp, key=key):
    n += 1
    ws.append(['NA%d' % n, 1, r, parts[r]['value'], '0402WGF4701TCE', 'UNI-ROYAL', 'C25900 (DNP)', '0402 (1005 metric)',
               'Resistor, 4.7 kOhm, 1 %, 1/16 W', side[r], 'Do not populate: the I2C pull-ups are on the FC (R15/R16). Fit only for a host without pull-ups'])
end = ws.max_row
for row in ws.iter_rows(min_row=1, max_row=end):
    for c in row:
        c.font = Font(name=FONT, size=10); c.alignment = Alignment(wrap_text=True, vertical='top')
ws['A1'].font = Font(name=FONT, size=13, bold=True)
ws['A2'].font = Font(name=FONT, size=9, italic=True, color='595959')
style_header(ws, H); style_header(ws, NA_H)
ws.cell(na0, 1).font = Font(name=FONT, size=11, bold=True)
for r in list(range(H + 1, last + 1)) + list(range(NA_H + 1, end + 1)):
    for c in ws[r]:
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for c in ws[tot]:
    c.font = Font(name=FONT, size=10, bold=True)
for col, w in zip('ABCDEFGHIJK', (6, 9, 22, 14, 24, 22, 20, 24, 40, 9, 60)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A5'

# ---------------------------------------------------------------- Costing
cs = wb.create_sheet('Costing')
cs.append(['ACwing Flow rev A - component costing (assembled parts, USD)'])
cs.append(['Blue = input prices; black = formulas. "Estimate" rows are typical distributor prices, not quotes. PCB, assembly fees and the cable are not included.'])
cs.append([])
ch = ['References', 'Qty / board', 'MPN', 'Source', 'Unit price @10 boards', 'Unit price @100 boards', 'Cost / board @10', 'Cost / board @100', 'Price basis']
cs.append(ch); CH = 4
for (k, sd), refs in rows:
    t = TAB[k]; r = cs.max_row + 1
    cs.append([', '.join(sorted(refs, key=key)), len(refs), t[0], t[2], t[7], t[8], '=B%d*E%d' % (r, r), '=B%d*F%d' % (r, r), t[9]])
cl = cs.max_row
cs.append([])
cs.append(['Total per board', '=SUM(B%d:B%d)' % (CH + 1, cl), '', '', '', '', '=SUM(G%d:G%d)' % (CH + 1, cl), '=SUM(H%d:H%d)' % (CH + 1, cl), ''])
ct = cs.max_row
u1 = next(r for r in range(CH + 1, cl + 1) if cs.cell(r, 1).value == 'U1')
cs.append(['Share of U1 (PMW3901)', '', '', '', '', '', '=G%d/G%d' % (u1, ct), '=H%d/H%d' % (u1, ct), 'The flow sensor is the largest single cost'])
for row in cs.iter_rows(min_row=1, max_row=cs.max_row):
    for c in row:
        c.font = Font(name=FONT, size=10); c.alignment = Alignment(wrap_text=True, vertical='top')
cs['A1'].font = Font(name=FONT, size=13, bold=True)
cs['A2'].font = Font(name=FONT, size=9, italic=True, color='595959')
style_header(cs, CH)
for r in range(CH + 1, cl + 1):
    for col in (5, 6):
        cs.cell(r, col).font = Font(name=FONT, size=10, color='0000FF'); cs.cell(r, col).number_format = '$#,##0.000'
    for col in (7, 8):
        cs.cell(r, col).number_format = '$#,##0.000'
    for c in cs[r]:
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for col in (7, 8):
    cs.cell(ct, col).number_format = '$#,##0.00'; cs.cell(ct + 1, col).number_format = '0.0%'
for c in cs[ct]:
    c.font = Font(name=FONT, size=10, bold=True)
for col, w in zip('ABCDEFGHI', (22, 10, 24, 24, 13, 13, 14, 14, 16)):
    cs.column_dimensions[col].width = w
wb.save(os.path.join(OUT, PROJECT + '_BoM.xlsx'))

# ---------------------------------------------------------------- JLCPCB BoM
jw = Workbook(); js = jw.active; js.title = 'Sheet1'
js.append(['Comment', 'Designator', 'Footprint', 'JLCPCB Part #'])
jg = {}
for (k, sd), refs in rows:
    t = TAB[k]
    g = jg.setdefault(t[0], [t, []]); g[1] += refs
for mpn, (t, refs) in sorted(jg.items(), key=lambda kv: key(sorted(kv[1][1], key=key)[0])):
    js.append([mpn, ','.join(sorted(refs, key=key)), t[6], t[2] if t[2].startswith('C') and t[2][1:].isdigit() else ''])
for row in js.iter_rows():
    for c in row:
        c.font = Font(name=FONT, size=10, bold=(c.row == 1))
for col, w in zip('ABCD', (26, 30, 34, 16)):
    js.column_dimensions[col].width = w
jw.save(os.path.join(OUT, PROJECT + '_BoM_JLCPCB.xlsx'))
print('BoM lines', len(rows), 'placements', sum(len(v) for v in groups.values()), '| JLC lines', len(jg), '| not assembled', n)
